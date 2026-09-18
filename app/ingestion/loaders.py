"""
Document loading module.

Each loader takes a file path and returns a list of `PageContent` objects
(one per page/section) so downstream chunking can retain page numbers.
Adding a new format only requires writing one function and registering
it in `EXTENSION_LOADERS`.
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, List

import fitz  # PyMuPDF
from docx import Document as DocxDocument

from app.exceptions import CorruptedFileError, EmptyDocumentError, UnsupportedFileTypeError

logger = logging.getLogger(__name__)


@dataclass
class PageContent:
    """Raw text extracted from a single page/section of a document."""

    page_number: int
    text: str


def _clean_text(text: str) -> str:
    """Normalize whitespace and strip control characters from raw text."""
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_pdf(path: Path) -> List[PageContent]:
    """Extract text from a PDF file page-by-page using PyMuPDF."""
    try:
        doc = fitz.open(path)
    except Exception as exc:  # PyMuPDF raises various low-level errors
        raise CorruptedFileError(f"Could not open PDF file: {path.name}") from exc

    pages: List[PageContent] = []
    try:
        for i, page in enumerate(doc, start=1):
            raw_text = page.get_text("text")
            cleaned = _clean_text(raw_text)
            if cleaned:
                pages.append(PageContent(page_number=i, text=cleaned))
    except Exception as exc:
        raise CorruptedFileError(f"Failed reading pages from PDF: {path.name}") from exc
    finally:
        doc.close()

    if not pages:
        raise EmptyDocumentError(f"No extractable text found in PDF: {path.name}")
    return pages


def load_docx(path: Path) -> List[PageContent]:
    """Extract text from a DOCX file.

    DOCX has no reliable page concept, so the whole document is treated
    as a single logical "page" (page_number=1). Paragraph and table text
    are both included.
    """
    try:
        doc = DocxDocument(path)
    except Exception as exc:
        raise CorruptedFileError(f"Could not open DOCX file: {path.name}") from exc

    parts: List[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)

    full_text = _clean_text("\n".join(parts))
    if not full_text:
        raise EmptyDocumentError(f"No extractable text found in DOCX: {path.name}")

    return [PageContent(page_number=1, text=full_text)]


def load_txt(path: Path) -> List[PageContent]:
    """Load a plain text file."""
    try:
        raw = path.read_text(encoding="utf-8", errors="ignore")
    except Exception as exc:
        raise CorruptedFileError(f"Could not read TXT file: {path.name}") from exc

    cleaned = _clean_text(raw)
    if not cleaned:
        raise EmptyDocumentError(f"No extractable text found in TXT file: {path.name}")

    return [PageContent(page_number=1, text=cleaned)]


# Register new formats here to extend support without touching callers.
EXTENSION_LOADERS: dict[str, Callable[[Path], List[PageContent]]] = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".txt": load_txt,
}


def load_document(path: Path) -> List[PageContent]:
    """Dispatch to the correct loader based on file extension."""
    ext = path.suffix.lower()
    loader = EXTENSION_LOADERS.get(ext)
    if loader is None:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}'. Supported types: "
            f"{', '.join(EXTENSION_LOADERS)}"
        )
    logger.info("Loading document %s with %s", path.name, loader.__name__)
    return loader(path)
